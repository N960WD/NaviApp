"""Run: python -m pytest tests   (from the exam-formatter folder)"""

import io
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import app  # noqa: E402
from exam_formatter import convert, parse_pdf  # noqa: E402

HERE = Path(__file__).parent


@pytest.fixture(scope="session", autouse=True)
def fixtures():
    for script in ("make_sample_pdf.py", "make_canvas_pdf.py"):
        subprocess.run([sys.executable, str(HERE / script)], check=True, capture_output=True)


def test_parses_numbered_exam():
    exam = parse_pdf(str(HERE / "sample_exam.pdf"))
    assert exam.title == "Biology 101 Midterm Exam"
    assert [q.number for q in exam.questions] == [1, 2, 3, 4, 5, 6]
    q1, q2 = exam.questions[:2]
    assert [o.label for o in q1.options] == ["A", "B", "C", "D"]
    assert [o.text for o in q2.options] == ["ATP", "DNA", "Glucose", "Lipids"]  # split from one line
    assert exam.questions[4].answer_blanks == [1, 2, 3]
    assert q1.points == "2 pts"


def test_parses_canvas_export():
    exam = parse_pdf(str(HERE / "canvas_exam.pdf"))
    assert len(exam.questions) == 8
    q1, q2 = exam.questions[:2]
    assert q1.answer_blanks == []  # "Blank 1:" field lines duplicate the inline blanks
    assert [o.text for o in q2.options] == ["Oxygen", "Nitrogen", "Argon", "Carbon dioxide"]
    assert q1.points == "1 pt"


def test_blank_markers_become_writing_spaces(tmp_path):
    out = tmp_path / "out.docx"
    convert(str(HERE / "sample_exam.pdf"), str(out))
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "Blank 1" not in text and "Blank 2" not in text
    assert " " * 24 in text


def test_upload_endpoint_returns_docx():
    client = app.test_client()
    with open(HERE / "sample_exam.pdf", "rb") as f:
        resp = client.post("/convert", data={"file": (f, "sample_exam.pdf")})
    assert resp.status_code == 200
    assert "sample_exam (formatted).docx" in resp.headers["Content-Disposition"]
    Document(io.BytesIO(resp.data))  # valid .docx

    bad = client.post("/convert", data={"file": (io.BytesIO(b"nope"), "x.pdf")})
    assert bad.status_code == 400
