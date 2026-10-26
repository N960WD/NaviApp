"""Render an Exam model as a clean, modern Word document."""

from __future__ import annotations

import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .parser import BLANK_RE, UNDERSCORES_RE, Exam, Question, Section

FONT = "Calibri"
NAVY = RGBColor(0x1F, 0x2A, 0x44)
ACCENT = RGBColor(0x2B, 0x6C, 0xB0)
MUTED = RGBColor(0x6B, 0x72, 0x80)
TEXT = RGBColor(0x1F, 0x23, 0x28)
ACCENT_HEX = "2B6CB0"
TINT_HEX = "EAF2FB"
RULE_HEX = "9AA5B1"

BLANK_WIDTH = 24  # non-breaking spaces per inline blank
WRITING_LINES = 4  # ruled lines under open-ended questions


# ---------------------------------------------------------------- xml helpers


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_borders(cell, **edges: tuple[str, int]) -> None:
    """edges: top/left/bottom/right=(hex color or 'nil', size in eighths of a point)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        color, size = edges.get(edge, ("nil", 0))
        if color == "nil":
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), color)
        borders.append(el)
    tc_pr.append(borders)


def _set_cell_margins(cell, top=100, bottom=100, left=180, right=180) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def _paragraph_bottom_rule(paragraph, color: str, size: int = 12, space: int = 6) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    bdr.append(bottom)
    p_pr.append(bdr)


def _run_shading(run, fill: str) -> None:
    r_pr = run._r.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    r_pr.append(shd)


def _run_underline(run, color: str) -> None:
    r_pr = run._r.get_or_add_rPr()
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    u.set(qn("w:color"), color)
    r_pr.append(u)


def _add_field(paragraph, instr: str, size: Pt, color: RGBColor) -> None:
    run = paragraph.add_run()
    run.font.size, run.font.color.rgb = size, color
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


def _no_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


# ---------------------------------------------------------------- styles


def _setup_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin, sec.bottom_margin = Inches(0.8), Inches(0.8)
    sec.header_distance = sec.footer_distance = Inches(0.4)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.space_before, pf.space_after = Pt(0), Pt(4)
    pf.line_spacing = 1.15
    return doc


def _content_width(doc) -> int:
    sec = doc.sections[0]
    return sec.page_width - sec.left_margin - sec.right_margin


# ---------------------------------------------------------------- blanks


def _add_blank(paragraph, number: int | None, width: int = BLANK_WIDTH) -> None:
    """A numbered tag followed by a clean underlined writing space."""
    if number is not None:
        tag = paragraph.add_run(f" {number} ")
        tag.font.size = Pt(8)
        tag.font.bold = True
        tag.font.color.rgb = ACCENT
        _run_shading(tag, TINT_HEX)
        tag.font.position = Pt(1)  # nudge up so it sits on the writing line's midline
        paragraph.add_run(" ")
    space = paragraph.add_run(" " * width)
    _run_underline(space, RULE_HEX)


def _add_rich_text(paragraph, text: str, size: Pt | None = None) -> int:
    """Write text, replacing 'Blank N' markers and ____ runs with writing spaces.

    Returns the number of blanks rendered.
    """
    pattern = re.compile(f"{BLANK_RE.pattern}|{UNDERSCORES_RE.pattern}", re.IGNORECASE)
    pos, count = 0, 0
    for m in pattern.finditer(text):
        before = text[pos:m.start()]
        if before:
            run = paragraph.add_run(before)
            if size:
                run.font.size = size
        number = int(m.group(1)) if m.group(1) else None
        if before and not before.endswith(" "):
            paragraph.add_run(" ")
        _add_blank(paragraph, number, BLANK_WIDTH if number else max(14, len(m.group(0))))
        nxt = text[m.end():m.end() + 1]
        if nxt and nxt.isalnum():
            paragraph.add_run(" ")
        pos = m.end()
        count += 1
    tail = text[pos:]
    if tail:
        run = paragraph.add_run(tail)
        if size:
            run.font.size = size
    return count


def _add_answer_line(doc, number: int, indent) -> None:
    """A full-width labelled response line for a blank listed on its own."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = indent
    pf.space_before, pf.space_after = Pt(8), Pt(2)
    pf.keep_with_next = True
    tag = p.add_run(f" {number} ")
    tag.font.size, tag.font.bold, tag.font.color.rgb = Pt(8), True, ACCENT
    _run_shading(tag, TINT_HEX)
    label = p.add_run("  Answer  ")
    label.font.size, label.font.color.rgb = Pt(9), MUTED
    p.paragraph_format.tab_stops.add_tab_stop(_content_width(doc), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES)
    line = p.add_run("\t")
    _run_underline(line, RULE_HEX)


# ---------------------------------------------------------------- sections


def _add_header_block(doc, exam: Exam) -> None:
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run(exam.title)
    run.font.size, run.font.bold, run.font.color.rgb = Pt(24), True, NAVY

    if exam.subtitle:
        sub = doc.add_paragraph()
        sub.paragraph_format.space_after = Pt(0)
        run = sub.add_run("  ·  ".join(exam.subtitle))
        run.font.size, run.font.color.rgb = Pt(11.5), MUTED

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(14)
    _paragraph_bottom_rule(rule, ACCENT_HEX, size=18, space=4)
    rule.paragraph_format.line_spacing = Pt(4)

    # Student info row: Name ______  Date ______  Score ____ / ____
    total = _total_points(exam)
    info = doc.add_paragraph()
    info.paragraph_format.space_after = Pt(16)
    width = _content_width(doc)
    stops = info.paragraph_format.tab_stops
    stops.add_tab_stop(int(width * 0.52), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.SPACES)
    stops.add_tab_stop(int(width * 0.56), WD_TAB_ALIGNMENT.LEFT)
    stops.add_tab_stop(int(width * 0.78), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.SPACES)
    stops.add_tab_stop(int(width * 0.82), WD_TAB_ALIGNMENT.LEFT)
    stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES)

    def label(text):
        r = info.add_run(text + "  ")
        r.font.size, r.font.bold, r.font.color.rgb = Pt(9), True, MUTED

    def line():
        r = info.add_run("\t")
        _run_underline(r, RULE_HEX)

    label("NAME")
    line()
    info.add_run("\t")
    label("DATE")
    line()
    info.add_run("\t")
    label("SCORE")
    line()
    if total:
        r = info.add_run(f" / {total}")
        r.font.size, r.font.color.rgb = Pt(10), MUTED

    if exam.instructions:
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        _set_cell_shading(cell, TINT_HEX)
        _set_cell_borders(cell, left=(ACCENT_HEX, 24))
        _set_cell_margins(cell, top=120, bottom=120, left=220, right=220)
        head = cell.paragraphs[0]
        r = head.add_run("INSTRUCTIONS")
        r.font.size, r.font.bold, r.font.color.rgb = Pt(9), True, ACCENT
        head.paragraph_format.space_after = Pt(3)
        for text in exam.instructions:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            _add_rich_text(p, text, size=Pt(10))
        doc.add_paragraph().paragraph_format.space_after = Pt(4)


def _add_section(doc, section: Section) -> None:
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(16), Pt(8)
    pf.keep_with_next = True
    _paragraph_bottom_rule(p, "D5DCE4", size=6, space=3)
    run = p.add_run(section.title.upper())
    run.font.size, run.font.bold, run.font.color.rgb = Pt(12), True, NAVY
    for text in section.paragraphs:
        sp = doc.add_paragraph()
        sp.paragraph_format.keep_with_next = True
        sp.paragraph_format.space_after = Pt(6)
        r_count = _add_rich_text(sp, text, size=Pt(10.5))
        if not r_count:
            for r in sp.runs:
                r.font.italic, r.font.color.rgb = True, MUTED


def _add_question(doc, q: Question) -> None:
    indent = Inches(0.42)
    paras = []

    stem_texts = q.paragraphs or [""]
    for i, text in enumerate(stem_texts):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = indent
        pf.space_after = Pt(4)
        pf.line_spacing = 1.35  # extra room so inline blanks can be written on
        if i == 0:
            pf.first_line_indent = -indent
            pf.space_before = Pt(12)
            pf.tab_stops.add_tab_stop(indent)
            num = p.add_run(f"{q.number}.")
            num.font.bold, num.font.color.rgb, num.font.size = True, ACCENT, Pt(12)
            p.add_run("\t")
        _add_rich_text(p, text)
        if i == len(stem_texts) - 1 and q.points:
            pts = p.add_run(f"   {q.points}")
            pts.font.size, pts.font.color.rgb, pts.font.italic = Pt(9), MUTED, True
        paras.append(p)

    for opt in q.options:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        opt_indent = indent + Inches(0.38)
        pf.left_indent = opt_indent
        pf.first_line_indent = -Inches(0.38)
        pf.space_after = Pt(3)
        pf.tab_stops.add_tab_stop(opt_indent)
        bubble = p.add_run("○ ")
        bubble.font.color.rgb, bubble.font.size = MUTED, Pt(11)
        lab = p.add_run(f"{opt.label}")
        lab.font.bold, lab.font.color.rgb, lab.font.size = True, ACCENT, Pt(10)
        p.add_run("\t")
        _add_rich_text(p, opt.text)
        paras.append(p)

    for n in q.answer_blanks:
        _add_answer_line(doc, n, indent)
        paras.append(doc.paragraphs[-1])

    # Open-ended questions get ruled writing lines.
    has_blank = any(BLANK_RE.search(t) or UNDERSCORES_RE.search(t) for t in q.paragraphs)
    if not q.options and not q.answer_blanks and not has_blank:
        for _ in range(WRITING_LINES):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = indent
            pf.space_before, pf.space_after = Pt(0), Pt(0)
            pf.line_spacing = Pt(26)
            pf.tab_stops.add_tab_stop(_content_width(doc), WD_TAB_ALIGNMENT.RIGHT)
            _run_underline(p.add_run("\t"), "C9D1DB")
            paras.append(p)

    # Keep each question on one page where possible.
    for p in paras[:-1]:
        p.paragraph_format.keep_with_next = True
    for p in paras:
        p.paragraph_format.keep_together = True
    if paras:
        paras[-1].paragraph_format.space_after = Pt(8)


def _add_footer(doc, exam: Exam) -> None:
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    width = _content_width(doc)
    style = doc.styles["Footer"]
    style.font.size, style.font.color.rgb = Pt(8.5), MUTED
    # The built-in Footer style has center/right stops at 3.25"/6.5"; clear them.
    for pos in (Inches(3.25), Inches(6.5)):
        p.paragraph_format.tab_stops.add_tab_stop(pos, WD_TAB_ALIGNMENT.CLEAR)
    p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(exam.title)
    r.font.size, r.font.color.rgb = Pt(8.5), MUTED
    p.add_run("\t")
    r = p.add_run("Page ")
    r.font.size, r.font.color.rgb = Pt(8.5), MUTED
    _add_field(p, "PAGE", Pt(8.5), MUTED)
    r = p.add_run(" of ")
    r.font.size, r.font.color.rgb = Pt(8.5), MUTED
    _add_field(p, "NUMPAGES", Pt(8.5), MUTED)


def _total_points(exam: Exam) -> str | None:
    vals = []
    for q in exam.questions:
        if not q.points:
            return None
        vals.append(float(q.points.split()[0]))
    if not vals:
        return None
    total = sum(vals)
    return str(int(total)) if total.is_integer() else f"{total:g}"


# ---------------------------------------------------------------- entry point


def build_docx(exam: Exam, out_path: str) -> str:
    doc = _setup_document()
    doc.core_properties.title = exam.title
    _add_header_block(doc, exam)
    for block in exam.blocks:
        if isinstance(block, Section):
            _add_section(doc, block)
        else:
            _add_question(doc, block)
    _add_footer(doc, exam)
    doc.save(out_path)
    return out_path
